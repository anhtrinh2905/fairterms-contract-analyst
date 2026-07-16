import pytest
from unittest.mock import MagicMock

from app.services.document_extractor import (
    DocumentKind,
    classify_pdf_kind,
    extract_docx_markdown,
    extract_pdf_markdown,
    has_meaningful_text,
    is_page_scanned,
    plain_text_to_markdown,
)


def test_has_meaningful_text_requires_letters():
    assert has_meaningful_text("12345") is False
    assert has_meaningful_text("Điều 1. Phạm vi điều chỉnh") is True


def test_is_page_scanned_uses_char_threshold():
    page = MagicMock()
    page.get_text.return_value = "   "
    assert is_page_scanned(page) is True
    page.get_text.return_value = "Điều 1. " * 20
    assert is_page_scanned(page) is False


def test_classify_pdf_kind_digital():
    document = MagicMock()
    document.page_count = 2
    page = MagicMock()
    page.get_text.return_value = "Điều 1. Nội dung hợp đồng dài đủ ký tự." * 5
    document.__getitem__.side_effect = lambda idx: page
    assert classify_pdf_kind(document) == DocumentKind.DIGITAL_PDF


def test_classify_pdf_kind_scanned():
    document = MagicMock()
    document.page_count = 2
    page = MagicMock()
    page.get_text.return_value = ""
    document.__getitem__.side_effect = lambda idx: page
    assert classify_pdf_kind(document) == DocumentKind.SCANNED_PDF


def test_plain_text_to_markdown_promotes_dieu_heading():
    markdown = plain_text_to_markdown("Điều 1. Phạm vi\nNội dung điều khoản.")
    assert "## Điều 1. Phạm vi" in markdown


def test_extract_pdf_markdown_from_digital_pdf():
    try:
        import fitz
    except ImportError:
        pytest.skip("PyMuPDF not installed")

    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Dieu 1. Pham vi dieu chinh hop dong thu nghiem.", fontname="helv")
    pdf_bytes = document.tobytes()
    document.close()

    markdown = extract_pdf_markdown(pdf_bytes)
    assert "Pham vi dieu chinh" in markdown


def test_extract_docx_markdown_includes_tables() -> None:
    try:
        from docx import Document
        from io import BytesIO
    except ImportError:
        pytest.skip("python-docx not installed")

    document = Document()
    document.add_paragraph("PHỤ LỤC HỢP ĐỒNG")
    document.add_paragraph("DANH MỤC THIẾT BỊ CĂN HỘ")
    table = document.add_table(rows=2, cols=4)
    table.rows[0].cells[0].text = "STT"
    table.rows[0].cells[1].text = "ĐỒ ĐẠC VÀ THIẾT BỊ"
    table.rows[0].cells[2].text = "SỐ LƯỢNG"
    table.rows[0].cells[3].text = "TÌNH TRẠNG"
    table.rows[1].cells[0].text = "1"
    table.rows[1].cells[1].text = "Điều hòa"
    table.rows[1].cells[2].text = "1"
    table.rows[1].cells[3].text = "Sử dụng tốt"

    buffer = BytesIO()
    document.save(buffer)
    markdown = extract_docx_markdown(buffer.getvalue())

    assert "DANH MỤC THIẾT BỊ" in markdown
    assert "Điều hòa" in markdown
    assert "| STT |" in markdown
    assert "Sử dụng tốt" in markdown
