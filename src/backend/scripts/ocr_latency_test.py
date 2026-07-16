"""Create sample uploads and measure /api/ocr/contract latency."""

from __future__ import annotations

import json
import sys
import time
from pathlib import Path

import httpx

FIXTURES = Path(__file__).resolve().parents[1] / "tests" / "fixtures" / "ocr_samples"
BACKEND_URL = "http://127.0.0.1:8010/api/ocr/contract"


def create_fixtures() -> dict[str, Path]:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}

    from docx import Document

    docx_path = FIXTURES / "sample_contract.docx"
    document = Document()
    document.add_heading("HỢP ĐỒNG THUÊ CĂN HỘ", 0)
    document.add_paragraph("Điều 1. Phạm vi điều chỉnh")
    document.add_paragraph(
        "Hợp đồng này quy định về việc thuê căn hộ chung cư giữa các bên."
    )
    document.add_paragraph("Điều 2. Thời hạn thuê")
    document.add_paragraph("Thời hạn thuê là 12 tháng kể từ ngày bàn giao căn hộ.")
    document.save(docx_path)
    paths["docx"] = docx_path

    import fitz

    pdf_path = FIXTURES / "sample_digital.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text(
        (72, 72),
        "HOP DONG THUE CAN HO\nDieu 1. Pham vi dieu chinh\nDieu 2. Thoi han thue 12 thang.",
        fontname="helv",
    )
    pdf.save(pdf_path)
    pdf.close()
    paths["digital_pdf"] = pdf_path

    doc_path = FIXTURES / "sample_contract.doc"
    if docx_path.exists():
        import shutil
        import subprocess

        soffice = shutil.which("soffice") or shutil.which("soffice.exe")
        for candidate in (
            Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
            Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
        ):
            if candidate.is_file():
                soffice = str(candidate)
                break
        if soffice:
            subprocess.run(
                [
                    soffice,
                    "--headless",
                    "--convert-to",
                    "doc",
                    "--outdir",
                    str(FIXTURES),
                    str(docx_path),
                ],
                check=True,
                capture_output=True,
                timeout=120,
            )
            converted = FIXTURES / "sample_contract.doc"
            if converted.is_file():
                paths["doc"] = converted

    return paths


def run_case(label: str, file_path: Path) -> dict:
    started = time.perf_counter()
    with file_path.open("rb") as handle:
        response = httpx.post(
            BACKEND_URL,
            files={"file": (file_path.name, handle, "application/octet-stream")},
            timeout=300.0,
        )
    elapsed_ms = int((time.perf_counter() - started) * 1000)
    result = {
        "label": label,
        "file": file_path.name,
        "http_status": response.status_code,
        "client_elapsed_ms": elapsed_ms,
    }
    if response.is_success:
        payload = response.json()
        result.update(
            {
                "processing_time_ms": payload.get("processing_time_ms"),
                "timings_ms": payload.get("timings_ms"),
                "markdown_chars": len(payload.get("markdown") or ""),
                "markdown_preview": (payload.get("markdown") or "")[:120],
            }
        )
    else:
        result["error"] = response.text[:300]
    return result


def main() -> int:
    try:
        health = httpx.get("http://127.0.0.1:8010/health", timeout=5.0)
        health.raise_for_status()
    except Exception as exc:
        print(f"Backend not reachable at :8010 — start uvicorn first. ({exc})")
        return 1

    paths = create_fixtures()
    print(f"Fixtures: {FIXTURES}\n")

    cases = [("docx", paths.get("docx")), ("digital_pdf", paths.get("digital_pdf"))]
    if "doc" in paths:
        cases.append(("doc", paths["doc"]))
    else:
        print("SKIP .doc — LibreOffice not installed (cannot build sample .doc)\n")

    for extra in sorted(FIXTURES.glob("*")):
        if extra.suffix.lower() in {".doc", ".docx", ".pdf"} and extra not in paths.values():
            cases.append((f"user_{extra.stem}", extra))

    results = []
    for label, path in cases:
        if path is None or not path.is_file():
            continue
        print(f"Testing {label} ({path.name})...")
        results.append(run_case(label, path))

    print("\n=== OCR latency report ===")
    print(json.dumps(results, ensure_ascii=False, indent=2))
    return 0 if all(item.get("http_status") == 200 for item in results) else 2


if __name__ == "__main__":
    raise SystemExit(main())
